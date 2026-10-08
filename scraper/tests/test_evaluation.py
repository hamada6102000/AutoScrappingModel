"""
Integration tests for the evaluation harness, driven through the real
management commands on local fixture pages with offline/scripted providers.
These verify the machinery only; they produce no research results.
"""

import csv
import io
import zipfile
from pathlib import Path

from django.core.management import call_command
from django.test import TestCase

from scraper.evaluation.config import ConfigError, resolve
from scraper.evaluation.protocol import EvaluationRunner, build_schedule
from scraper.llm.base import LLMProvider, RateLimitExhausted
from scraper.llm.heuristic import HeuristicProvider
from scraper.models import (
    Experiment, GroundTruth, GTSource, LLMCallLog, RulePurpose, ScrapeRun, Snapshot,
    SnapshotKind, XPathRule,
)
from scraper.services import snapshots as snapstore

from .helpers import SELFTEST_CONFIG, SELFTEST_GT, TempDataDirMixin


def quiet(*args, **kwargs):
    out = io.StringIO()
    call_command(*args, stdout=out, stderr=out, **kwargs)
    return out.getvalue()


class FlakyProvider(LLMProvider):
    """Heuristic answers, but the quota 'runs out' after N calls."""

    name = "groq"
    model = "flaky-model"

    def __init__(self, ok_calls):
        self.left = ok_calls
        self.inner = HeuristicProvider()

    def generate_xpaths(self, pruned_html, fields, failure_history=None):
        if self.left <= 0:
            raise RateLimitExhausted("daily tokens exhausted")
        self.left -= 1
        return self.inner.generate_xpaths(pruned_html, fields, failure_history)


class BrokenProvider(LLMProvider):
    """The provider cannot be reached at all (e.g. prompt larger than the TPM limit)."""

    name = "groq"
    model = "broken-model"

    def __init__(self, exc):
        self.exc = exc

    def generate_xpaths(self, pruned_html, fields, failure_history=None):
        raise self.exc


class HarnessTests(TempDataDirMixin, TestCase):
    def prepare(self, xlsx=False):
        quiet("eval_init", config=str(SELFTEST_CONFIG))
        quiet("eval_snapshot")
        quiet("eval_gt_template")
        exp = Experiment.objects.get(name="selftest")
        base = snapstore.experiment_dir(exp) / "ground_truth_baseline"
        if xlsx:
            from openpyxl import load_workbook

            wb = load_workbook(base.with_suffix(".xlsx"))
            ws = wb["ground_truth"]
            header = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2):
                cells = dict(zip(header, row))
                cells["value"].value = SELFTEST_GT[cells["site_key"].value][cells["field"].value]
            # Simulate Excel turning a typed number into a numeric cell.
            for row in ws.iter_rows(min_row=2):
                cells = dict(zip(header, row))
                if cells["site_key"].value == "shop-product" and cells["field"].value == "rating":
                    cells["value"].value = 4.3
            filled = Path(self._tmp) / "filled.xlsx"
            wb.save(filled)
        else:
            with open(base.with_suffix(".csv"), encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.DictReader(fh))
            for r in rows:
                r["value"] = SELFTEST_GT[r["site_key"]][r["field"]]
            filled = Path(self._tmp) / "filled.csv"
            with open(filled, "w", encoding="utf-8-sig", newline="") as fh:
                w = csv.DictWriter(fh, fieldnames=list(rows[0]))
                w.writeheader()
                w.writerows(rows)
        out = quiet("eval_gt_import", file=str(filled), verified_by="tester")
        self.assertIn("All baseline labels present", out)
        return exp

    def test_full_workflow(self):
        exp = self.prepare()
        # 5 sites x 2 conditions x 2 repetitions
        self.assertEqual(exp.targets.count(), 20)
        self.assertEqual(Snapshot.objects.filter(experiment=exp, kind=SnapshotKind.BASELINE).count(), 5)
        for snap in Snapshot.objects.filter(experiment=exp):
            self.assertTrue((snapstore.data_dir() / snap.path).exists())
            self.assertTrue(snapstore.label_view_path(snap).exists())

        out = quiet("eval_run")
        self.assertIn("Done this session: 180 runs", out)
        self.assertEqual(ScrapeRun.objects.filter(experiment=exp, completed=True).count(), 180)

        # Static condition: identical t0 rules, never calls the LLM.
        self.assertFalse(LLMCallLog.objects.filter(run__experiment=exp, run__condition="static").exists())
        for rule in XPathRule.objects.filter(field__target__experiment=exp,
                                             field__target__condition="static"):
            self.assertEqual(rule.purpose, RulePurpose.COPIED)
            self.assertEqual(rule.xpath, rule.copied_from.xpath)

        # Both conditions saw byte-identical snapshots at every step.
        pairs = {}
        for run in ScrapeRun.objects.filter(experiment=exp):
            pairs.setdefault((run.target.site_key, run.repetition, run.step), set()).add(run.snapshot_id)
        self.assertTrue(all(len(v) == 1 for v in pairs.values()))

        # Simulated snapshots carry ground truth by construction.
        sim = Snapshot.objects.filter(experiment=exp, kind=SnapshotKind.SIMULATED)
        self.assertEqual(sim.count(), 5 * 2 * 8)
        sources = set(GroundTruth.objects.filter(snapshot__in=sim).values_list("source", flat=True))
        self.assertTrue(sources <= {GTSource.DERIVED, GTSource.PERTURBED})
        self.assertFalse(GroundTruth.objects.filter(snapshot__in=sim, notes__startswith="exclude").exists())

        # Resuming does nothing.
        out = quiet("eval_run")
        self.assertIn("0 runs", out)
        self.assertEqual(ScrapeRun.objects.filter(experiment=exp).count(), 180)

        out = quiet("eval_report")
        self.assertIn("Initial extraction accuracy", out)
        results = snapstore.experiment_dir(exp) / "results"
        for name in ("summary.json", "report.md", "tables/table_5_1.csv", "tables/detection_rq1.csv",
                     "tables/f1_by_step.csv", "tables/healing_events_rq4.csv"):
            self.assertTrue((results / name).exists(), name)

        zip_path = Path(self._tmp) / "bundle.zip"
        quiet("export_results", out=str(zip_path))
        names = set(zipfile.ZipFile(zip_path).namelist())
        for name in ("experiment.json", "tables/runs.csv", "tables/records.csv", "tables/llm_calls.csv",
                     "tables/ground_truth.csv", "results/report.md"):
            self.assertIn(name, names)
        self.assertTrue(any(n.endswith("baseline_baseline.html") for n in names))

    def test_xlsx_ground_truth_with_numeric_cells(self):
        exp = self.prepare(xlsx=True)
        snap = Snapshot.objects.get(experiment=exp, site_key="shop-product", kind=SnapshotKind.BASELINE)
        self.assertEqual(GroundTruth.objects.get(snapshot=snap, field_name="rating").value, "4.3")
        listing = Snapshot.objects.get(experiment=exp, site_key="category-listing", kind=SnapshotKind.BASELINE)
        self.assertEqual(len(GroundTruth.objects.get(snapshot=listing, field_name="title").values), 12)

    def test_quota_stop_rolls_back_and_resumes(self):
        exp = self.prepare()
        runner = EvaluationRunner(exp, write=lambda m: None, provider=FlakyProvider(ok_calls=3),
                                  sites=["shop-product"])
        stats = runner.run()
        self.assertIn("quota", stats["stopped"])
        self.assertFalse(ScrapeRun.objects.filter(experiment=exp, completed=False).exists())
        done = ScrapeRun.objects.filter(experiment=exp).count()
        # Every stored LLM call belongs to a completed run (the interrupted one was rolled back).
        self.assertFalse(LLMCallLog.objects.filter(run__experiment=exp, run__completed=False).exists())

        stats = EvaluationRunner(exp, write=lambda m: None, provider=FlakyProvider(ok_calls=10_000),
                                 sites=["shop-product"]).run()
        self.assertNotIn("stopped", stats)
        self.assertEqual(stats["skipped"], done)
        self.assertEqual(ScrapeRun.objects.filter(experiment=exp, completed=True).count(), 2 * 2 * 9)

    def test_provider_errors_stop_instead_of_counting_as_failures(self):
        from scraper.llm.base import RequestTooLarge
        exp = self.prepare()
        stats = EvaluationRunner(exp, write=lambda m: None, sites=["shop-product"],
                                 provider=BrokenProvider(RequestTooLarge("too large"))).run()
        self.assertIn("provider error", stats["stopped"])
        self.assertFalse(stats["retryable"])
        self.assertFalse(ScrapeRun.objects.filter(experiment=exp).exists())
        stats = EvaluationRunner(exp, write=lambda m: None, sites=["shop-product"],
                                 provider=BrokenProvider(ConnectionError("reset"))).run()
        self.assertTrue(stats["retryable"])
        self.assertFalse(LLMCallLog.objects.filter(run__experiment=exp).exists())

    def test_llm_call_budget(self):
        exp = self.prepare()
        out = quiet("eval_run", max_llm_calls=1, sites="shop-product")
        self.assertIn("--max-llm-calls budget", out)

    def test_readiness_requires_labels(self):
        quiet("eval_init", config=str(SELFTEST_CONFIG))
        quiet("eval_snapshot", sites="shop-product")
        exp = Experiment.objects.get(name="selftest")
        problems = EvaluationRunner(exp, write=lambda m: None).readiness()
        self.assertTrue(any("ground truth missing" in p for p in problems))
        self.assertTrue(any("no baseline snapshot" in p for p in problems))

    def test_external_baseline_import(self):
        exp = self.prepare()
        snap = Snapshot.objects.get(experiment=exp, site_key="shop-product", kind=SnapshotKind.BASELINE)
        preds = Path(self._tmp) / "preds.csv"
        with open(preds, "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["snapshot_id", "field", "value"])
            w.writerow([snap.pk, "name", "Nova Smart Watch"])
            w.writerow([snap.pk, "price", "$189.00"])  # wrong
        quiet("eval_import_predictions", name="autoscraper", file=str(preds))
        quiet("eval_run", sites="shop-product")
        quiet("eval_report")
        from scraper.evaluation.report import build

        rows = {b["system"]: b for b in build(exp)["baselines"]}
        self.assertIn("autoscraper", rows)
        self.assertEqual(rows["autoscraper"]["tp"], 1)
        self.assertIn("self_healing (same snapshots)", rows)

    def test_sensitivity_sets_aside_by_audit_rule_only(self):
        from scraper.evaluation import report as report_mod
        from scraper.evaluation import sensitivity as sens

        exp = self.prepare()
        quiet("eval_run")
        sites = [s["key"] for s in exp.config["sites"]]
        # audit rows (normally from run_audit): one site with 2 of 4 values missing, one with 1 of 4
        rows = []
        for i, key in enumerate(sites):
            missing = 2 if i == 0 else (1 if i == 1 else 0)
            rows.append({"site": key, "fields": [{"field": f"f{j}", "status": "not-found" if j < missing else "found",
                                                  "excluded": False} for j in range(4)]
                         + [{"field": "x", "status": "not-found", "excluded": True}]})
        rep = sens.build(exp, 0.25, audit_rows=rows)
        self.assertEqual(rep["sites_set_aside"], sorted(sites[:2]))  # 50% and 25% are >= 25%
        self.assertAlmostEqual(rep["site_shares"][sites[1]]["share"], 0.25)
        rep = sens.build(exp, 0.3, audit_rows=rows)
        self.assertEqual(rep["sites_set_aside"], [sites[0]])
        full, sub = rep["all_sites"], rep["without_set_aside"]
        self.assertEqual(full["sites"], sorted(sites))
        self.assertEqual(sub["sites"], sorted(sites[1:]))
        for t in sub["paired_tests"]:
            self.assertLessEqual(t["sites"], len(sites) - 1)
        # the all-sites view reproduces the primary report exactly
        primary = {m["condition"]: m for m in report_mod.build(exp)["modules"]}
        for m in full["modules"]:
            for k in ("initial_f1", "changed_f1", "overall_f1", "healing_success_rate",
                      "false_alarms_per_monitoring_run"):
                self.assertEqual(m[k], primary[m["condition"]][k], k)
        self.assertIn("Post hoc", sens.markdown(rep))
        out = quiet("eval_sensitivity", min_not_found_share=0.25)
        self.assertIn("Post hoc sensitivity analysis", out)
        results = snapstore.experiment_dir(exp) / "results"
        self.assertTrue((results / "sensitivity.md").exists() and (results / "sensitivity.json").exists())

    def test_schedule_is_deterministic_and_complete(self):
        cfg = resolve({"experiment": {"name": "x"}, "sites": [{"key": "a", "url": "https://a.test",
                                                               "fields": [{"name": "t"}]}]})
        e = cfg["experiment"]
        self.assertEqual(build_schedule(e, "a", 1), build_schedule(e, "a", 1))
        self.assertNotEqual([s["seed"] for s in build_schedule(e, "a", 1)],
                            [s["seed"] for s in build_schedule(e, "a", 2)])
        self.assertEqual(sorted(s["op"] for s in build_schedule(e, "a", 1)),
                         sorted(s["op"] for s in e["protocol"]["steps"]))


class ConfigTests(TestCase):
    def test_reports_all_problems_at_once(self):
        with self.assertRaises(ConfigError) as ctx:
            resolve({"experiment": {"name": "", "repetitions": 0, "llm": {"provider": "gpt9"}},
                     "sites": [{"key": "Bad Key!", "url": "ftp://x", "category": "toys",
                                "fields": [{"name": "p", "type": "money", "cardinality": "lots"}]}]})
        text = "\n".join(ctx.exception.problems)
        for fragment in ("experiment.name", "provider", "repetitions", "key must be", "url must",
                         "category", "type must", "cardinality"):
            self.assertIn(fragment, text)

    def test_selftest_config_is_valid(self):
        from scraper.evaluation.config import load_config

        cfg = load_config(SELFTEST_CONFIG)
        self.assertEqual(len(cfg["sites"]), 5)
        self.assertTrue(Path(cfg["sites"][0]["html_file"]).is_absolute())

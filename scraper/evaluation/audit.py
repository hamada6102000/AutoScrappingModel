"""
Configuration audit on the saved snapshots (run after eval_snapshot, before
labelling and before any LLM call).

For every site: was the page fetched, with which render mode, is it a real
content page (expected texts present), robots.txt, token sizes, and how big
the actual LLM prompt is. For every field: is the value the annotator recorded
present in the saved page — in the visible content, in an attribute, or only
in the page-metadata block (meta tags / JSON-LD) — and would the LLM see it in
the (possibly truncated) prompt.

It also writes a ground-truth sheet pre-filled with the recorded values that
were found literally in the snapshot (and abridged texts/lists completed from
the page element that holds them). Every pre-filled row says how it was
obtained; rows that could not be matched are left empty for a person to label.
Nothing is filled from the system's own output.
"""

from __future__ import annotations

import csv
from pathlib import Path

from django.conf import settings

from ..llm.base import FieldSpec
from ..llm.prompts import build_prompt
from ..services import snapshots as snapstore
from ..services import tokens as tok
from .dataset import ELLIPSIS, Located, PageIndex, locate, probes
from .gt_io import ABSENT, COLUMNS, INSTRUCTIONS
from .protocol import baseline_snapshot

NEEDS_LABEL = ("not-found", "not-literal", "list-partial")
PREFILL_NOTE = {
    "found": "prefilled: recorded value found in the snapshot",
    "found-partial": "prefilled: recorded value found inside a longer page text",
    "found-variant": "prefilled: recorded value found without its note",
    "prefix": "prefilled: abridged recorded text completed from the page element starting with it",
    "list-complete": "prefilled: all recorded items found in the snapshot",
    "list-completed": "prefilled: abridged recorded list completed from the page container of its items",
    "boolean": "prefilled: recorded yes/no value; check it on the page",
}


def max_input_chars(exp, override: int | None = None) -> int:
    if override:
        return int(override)
    llm = exp.config["experiment"].get("llm") or {}
    return int(llm.get("max_input_chars") or settings.SCRAPER.get("LLM_MAX_INPUT_CHARS", 60000))


def _specs(site) -> list[FieldSpec]:
    return [FieldSpec(name=f["name"], description=f["description"], field_type=f["type"],
                      cardinality=f["cardinality"]) for f in site["fields"]]


def _value_for_view(loc: Located):
    if isinstance(loc.value, list):
        return loc.value[0] if loc.value else None
    if isinstance(loc.value, str) and loc.status == "prefix":
        return loc.value[:60] + ELLIPSIS
    return loc.value


def audit_site(exp, site: dict, cap: int | None = None) -> dict:
    snap = baseline_snapshot(exp, site["key"])
    row = {"site": site["key"], "name": site["name"], "category": site["category"],
           "url": site["url"], "snapshot": bool(snap), "fields": []}
    all_fields = [dict(f, excluded=False) for f in site["fields"]] + \
                 [dict(f, excluded=True) for f in site.get("excluded_fields") or []]
    if snap is None:
        for f in all_fields:
            row["fields"].append({"field": f["name"], "status": "no-snapshot", "excluded": f["excluded"]})
        return row
    meta = snap.fetch_meta or {}
    check = meta.get("check") or {}
    page = snapstore.pruned_page(snap, site["prune_keep"])
    idx = PageIndex(page.clean_html, snap.url)
    cap = max_input_chars(exp, cap)
    prompt, pmeta = build_prompt(page.llm_view, _specs(site), max_html_chars=cap)
    view_html = page.llm_view[:cap] if cap else page.llm_view
    view_idx = PageIndex(view_html, snap.url)
    row.update({
        "http_status": snap.http_status,
        "render": meta.get("render_chosen") or meta.get("render") or ("file" if meta.get("source") == "html_file" else "static"),
        "render_candidates": meta.get("render_candidates") or {},
        "render_errors": meta.get("render_errors") or {},
        "robots": snap.robots_allowed,
        "suspected_block": bool(meta.get("suspected_block")),
        "content_ok": check.get("content_ok", True),
        "expect": f"{check.get('expect_found', '?')}/{check.get('expect_total', '?')}",
        "tokenizer": snap.tokenizer,
        "raw_tokens": snap.raw_tokens, "clean_tokens": snap.clean_tokens,
        "llm_view_tokens": snap.llm_view_tokens,
        "prompt_tokens": tok.count_tokens(prompt),
        "prompt_truncated": pmeta["input_truncated"],
        "has_metadata": idx.meta_root is not None,
        "labeling_view": str(snapstore.label_view_path(snap)),
        "snapshot_id": snap.pk,
    })
    for f in all_fields:
        loc = locate(f, idx)
        in_prompt = None
        probe_value = _value_for_view(loc)
        if probe_value and loc.status not in NEEDS_LABEL[:2]:
            in_prompt = bool(view_idx.where(probe_value, f["type"] if not probe_value.endswith(ELLIPSIS) else "string"))
        row["fields"].append({
            "field": f["name"], "type": f["type"], "cardinality": f["cardinality"],
            "required": f["required"], "dataset_source": f.get("source", ""),
            "dynamic": f.get("dynamic", False), "dataset_value": f.get("dataset_value", ""),
            "status": loc.status, "found_in": "+".join(loc.where), "in_prompt": in_prompt,
            "value": loc.value, "detail": loc.detail, "excluded": f["excluded"],
            "exclude_reason": f.get("exclude_reason", ""), "description": f["description"],
            "probe": probes(f).kind,
        })
    return row


def run_audit(exp, cap: int | None = None) -> list[dict]:
    return [audit_site(exp, s, cap) for s in exp.config["sites"]]


# ---------------------------------------------------------------------------
# Outputs
# ---------------------------------------------------------------------------
def _fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        return "\n".join(v)
    return str(v)


def write_outputs(exp, rows: list[dict], out_dir: Path, cap: int | None = None) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = {}

    site_cols = ["site", "name", "category", "snapshot", "http_status", "render", "robots",
                 "suspected_block", "content_ok", "expect", "raw_tokens", "clean_tokens",
                 "llm_view_tokens", "prompt_tokens", "prompt_truncated", "has_metadata",
                 "n_fields", "found", "needs_label", "only_metadata", "not_in_prompt", "url"]
    field_cols = ["site", "field", "type", "cardinality", "required", "excluded", "dataset_source",
                  "dynamic", "status", "found_in", "in_prompt", "dataset_value", "value", "detail"]
    site_rows = []
    for r in rows:
        fs = [f for f in r["fields"] if not f.get("excluded")]
        site_rows.append({
            **{k: r.get(k, "") for k in site_cols},
            "n_fields": len(fs),
            "found": sum(1 for f in fs if f["status"] not in NEEDS_LABEL and f["status"] != "no-snapshot"),
            "needs_label": sum(1 for f in fs if f["status"] in NEEDS_LABEL),
            "only_metadata": sum(1 for f in fs if f.get("found_in") == "metadata"),
            "not_in_prompt": sum(1 for f in fs if f.get("in_prompt") is False),
        })
    paths["sites_csv"] = out_dir / "audit_sites.csv"
    with open(paths["sites_csv"], "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=site_cols)
        w.writeheader()
        w.writerows(site_rows)
    paths["fields_csv"] = out_dir / "audit_fields.csv"
    with open(paths["fields_csv"], "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=field_cols)
        w.writeheader()
        for r in rows:
            for f in r["fields"]:
                w.writerow({**{k: f.get(k, "") for k in field_cols}, "site": r["site"],
                            "value": _fmt_value(f.get("value"))[:2000]})
    paths["markdown"] = out_dir / "audit.md"
    paths["markdown"].write_text(markdown(exp, rows, site_rows, cap), encoding="utf-8")
    paths["gt_sheet"] = write_prefilled_sheet(exp, rows, out_dir / "ground_truth_prefilled.xlsx")
    return paths


def markdown(exp, rows, site_rows, cap: int | None = None) -> str:
    cap = max_input_chars(exp, cap)
    L = [f"# Configuration audit — experiment `{exp.name}`", "",
         f"Snapshots checked against the annotated dataset. Prompt HTML cap: {cap} characters.",
         "", "## Sites", "",
         "| site | HTTP | render | robots | content check | LLM view tok | prompt tok | prompt cut | fields | found | to label | metadata only | not in prompt |",
         "|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for s in site_rows:
        L.append(f"| {s['site']} | {s['http_status']} | {s['render']} | {s['robots']} | "
                 f"{'ok' if s['content_ok'] else 'FAILED'} ({s['expect']}) | {s['llm_view_tokens']} | "
                 f"{s['prompt_tokens']} | {'yes' if s['prompt_truncated'] else 'no'} | {s['n_fields']} | "
                 f"{s['found']} | {s['needs_label']} | {s['only_metadata']} | {s['not_in_prompt']} |")
    L += ["", "## Fields that need attention", ""]
    for r in rows:
        issues = [f for f in r["fields"] if not f.get("excluded") and
                  (f["status"] in NEEDS_LABEL or f.get("in_prompt") is False or f["status"] == "no-snapshot")]
        if not issues:
            continue
        L.append(f"### {r['site']} — {r['name']}")
        for f in issues:
            where = f" (found in {f['found_in']})" if f.get("found_in") else ""
            prompt = " — NOT in the LLM prompt" if f.get("in_prompt") is False else ""
            L.append(f"- `{f['field']}` [{f.get('type', '')}, {f.get('dataset_source', '')}]: "
                     f"**{f['status']}**{where}{prompt}. {f.get('detail', '')}")
        L.append("")
    excluded = [(r["site"], f) for r in rows for f in r["fields"] if f.get("excluded")]
    if excluded:
        L += ["## Excluded fields", ""]
        for site, f in excluded:
            L.append(f"- {site}.`{f['field']}`: {f.get('exclude_reason', '')}")
    return "\n".join(L) + "\n"


def write_prefilled_sheet(exp, rows, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    cols = COLUMNS + ["prefill_status", "found_in", "in_prompt", "dataset_value", "dataset_source", "dynamic"]
    wb = Workbook()
    ws = wb.active
    ws.title = "ground_truth"
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    review = PatternFill("solid", fgColor="FCE4D6")
    editable = PatternFill("solid", fgColor="FFF8D6")
    for r in rows:
        if not r.get("snapshot"):
            continue
        for f in r["fields"]:
            if f.get("excluded"):
                continue
            prefill = f["status"] not in NEEDS_LABEL
            value = _fmt_value(f.get("value")) if prefill else ""
            if f["cardinality"] == "many" and isinstance(f.get("value"), list) and prefill:
                value = "\n".join(f["value"])
            ws.append([
                r["snapshot_id"], r["site"], "baseline", f["field"], f["type"], f["cardinality"],
                "yes" if f["required"] else "no", f["description"], value, "",
                (PREFILL_NOTE.get(f["status"], "") + (f"; {f['detail']}" if f.get("detail") else "")) if prefill else "",
                r["labeling_view"], f["status"], f.get("found_in", ""),
                "" if f.get("in_prompt") is None else ("yes" if f["in_prompt"] else "NO"),
                f.get("dataset_value", ""), f.get("dataset_source", ""), "yes" if f.get("dynamic") else "no",
            ])
            rown = ws.max_row
            for c in range(1, len(cols) + 1):
                cell = ws.cell(row=rown, column=c)
                cell.number_format = "@"
                cell.alignment = Alignment(wrap_text=cols[c - 1] in ("value", "dataset_value", "notes"),
                                           vertical="top")
            for name in ("value", "verified_by", "notes"):
                ws.cell(row=rown, column=cols.index(name) + 1).fill = editable
            if not prefill or f["status"] in ("boolean", "prefix", "list-completed", "found-variant",
                                               "found-partial"):
                ws.cell(row=rown, column=cols.index("prefill_status") + 1).fill = review
    widths = {"value": 50, "notes": 40, "description": 30, "dataset_value": 40, "labeling_view": 30}
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(col, 13)
    ws.freeze_panes = "E2"
    info = wb.create_sheet("README", 0)
    lines = [
        "PRE-FILLED GROUND-TRUTH SHEET",
        "",
        "Rows are pre-filled ONLY where the value you recorded in the dataset was found in the saved snapshot",
        "(column prefill_status). Please check every row against its labeling view, then:",
        "  - found / list-complete: the recorded value is on the saved page; confirm it.",
        "  - found-partial: the value is inside a longer text (e.g. '22' in 'In stock (22 available)'); confirm.",
        "  - found-variant: found after removing the note in parentheses; confirm the form is right.",
        "  - prefix / list-completed: your value was abridged; the full text/list was read from the page",
        "    element that holds it — check it is the right element and complete.",
        "  - boolean: yes/no value from the dataset — check it on the page.",
        "  - not-found / not-literal / list-partial: EMPTY — type the value as it appears in the labeling view,",
        f"    or {ABSENT} if the field is not on the page.",
        "Put your initials in verified_by on each row you checked (or pass --verified-by when importing).",
        "",
    ] + INSTRUCTIONS
    for line in lines:
        info.append([line])
    info.column_dimensions["A"].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path

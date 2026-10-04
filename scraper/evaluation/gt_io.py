"""
Ground-truth templates (Excel + CSV) and import.

The template is intentionally NOT pre-filled with system output: labels must
be independent of the system under test. Re-exporting after an import shows
the values already entered so they can be corrected.

Conventions in the 'value' column:
  * single field  -> the value exactly as shown in the labeling view
  * list field    -> one item per line (Alt+Enter in Excel) or items separated by ||
  * field not on the page -> <ABSENT>
  * empty cell    -> not labelled yet (the run refuses to start)
"""

from __future__ import annotations

import csv
import datetime as dt
import json
from pathlib import Path

from ..models import GroundTruth, GTSource, Snapshot, SnapshotKind
from ..services import snapshots as snapstore
from .config import site_by_key

ABSENT = "<ABSENT>"
COLUMNS = [
    "snapshot_id", "site_key", "snapshot_label", "field", "type", "cardinality",
    "required", "description", "value", "verified_by", "notes", "labeling_view",
]
INSTRUCTIONS = [
    "HOW TO FILL THIS SHEET",
    "",
    "1. Open the file in the 'labeling_view' column in your browser. It is the saved page with scripts removed,",
    "   i.e. exactly what the scraper sees. Do NOT use the live website: prices and text may have changed.",
    "2. In 'value', type or paste each field's value exactly as it appears there.",
    "   - List fields: one item per line (Alt+Enter inside the cell), or separate items with ||",
    f"   - If the field does not exist on that page, write {ABSENT}",
    "   - Single field with two acceptable forms (e.g. one date shown in two time zones): A || B",
    "3. Put your name/initials in 'verified_by'. Use 'notes' for anything ambiguous.",
    "4. Leave every other column unchanged. Save, then run:",
    "   python manage.py eval_gt_import --experiment <name> --file <this file>",
    "",
    "Empty 'value' cells mean 'not labelled yet'; the evaluation will not start until all are filled.",
]


def _snapshots(exp, kind: str, label: str | None):
    qs = Snapshot.objects.filter(experiment=exp, kind=kind)
    if label:
        qs = qs.filter(label=label)
    return list(qs.order_by("site_key", "fetched_at"))


def template_rows(exp, kind: str = SnapshotKind.BASELINE, label: str | None = "baseline") -> list[dict]:
    rows = []
    for snap in _snapshots(exp, kind, label):
        try:
            site = site_by_key(exp.config, snap.site_key)
        except KeyError:
            continue
        existing = {g.field_name: g for g in snap.ground_truth.all()}
        for f in site["fields"]:
            g = existing.get(f["name"])
            value = ""
            if g is not None:
                if g.notes == "absent":
                    value = ABSENT
                elif f["cardinality"] == "many":
                    value = "\n".join(g.values)
                else:
                    value = g.display_value
            rows.append({
                "snapshot_id": snap.pk,
                "site_key": snap.site_key,
                "snapshot_label": snap.label,
                "field": f["name"],
                "type": f["type"],
                "cardinality": f["cardinality"],
                "required": "yes" if f["required"] else "no",
                "description": f["description"],
                "value": value,
                "verified_by": g.verified_by if g else "",
                "notes": "" if (g is None or g.notes == "absent") else g.notes,
                "labeling_view": str(snapstore.label_view_path(snap)),
            })
    return rows


def write_template(exp, out_base: Path, kind: str = SnapshotKind.BASELINE,
                   label: str | None = "baseline") -> list[Path]:
    rows = template_rows(exp, kind, label)
    out_base.parent.mkdir(parents=True, exist_ok=True)
    written = []

    csv_path = out_base.with_suffix(".csv")
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    written.append(csv_path)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return written

    wb = Workbook()
    ws = wb.active
    ws.title = "ground_truth"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in rows:
        ws.append([row[c] for c in COLUMNS])
    value_col = COLUMNS.index("value") + 1
    editable = {COLUMNS.index(c) + 1 for c in ("value", "verified_by", "notes")}
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.number_format = "@"  # keep everything as text (no dates / lost zeros)
            cell.alignment = Alignment(wrap_text=(c == value_col), vertical="top")
            if c in editable:
                cell.fill = PatternFill("solid", fgColor="FFF8D6")
    widths = {"snapshot_id": 11, "site_key": 18, "snapshot_label": 14, "field": 16, "type": 9,
              "cardinality": 11, "required": 9, "description": 34, "value": 42,
              "verified_by": 13, "notes": 24, "labeling_view": 60}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(col, 14)
    ws.freeze_panes = "A2"

    info = wb.create_sheet("README", 0)
    for line in INSTRUCTIONS:
        info.append([line])
    info.column_dimensions["A"].width = 120

    if kind == SnapshotKind.NATURAL:
        nat = wb.create_sheet("natural_changes")
        nat.append(["snapshot_id", "site_key", "snapshot_label", "layout_changed (yes/no)", "notes"])
        for snap in _snapshots(exp, kind, label):
            known = (snap.mutation or {}).get("layout_changed")
            nat.append([snap.pk, snap.site_key, snap.label,
                        "" if known is None else ("yes" if known else "no"), ""])

    xlsx_path = out_base.with_suffix(".xlsx")
    wb.save(xlsx_path)
    written.append(xlsx_path)
    return written


# ---------------------------------------------------------------------------
def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    return str(value)


def read_rows(path: Path) -> tuple[list[dict], list[dict]]:
    """Return (ground_truth_rows, natural_change_rows)."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        ws = wb["ground_truth"] if "ground_truth" in wb.sheetnames else wb.active
        values = list(ws.iter_rows(values_only=True))
        header = [_cell_text(h).strip() for h in values[0]] if values else []
        rows = [dict(zip(header, (_cell_text(v) for v in row))) for row in values[1:]]
        natural = []
        if "natural_changes" in wb.sheetnames:
            nv = list(wb["natural_changes"].iter_rows(values_only=True))
            nh = ["snapshot_id", "site_key", "snapshot_label", "layout_changed", "notes"]
            natural = [dict(zip(nh, (_cell_text(v) for v in row))) for row in nv[1:]]
        return rows, natural
    for encoding in ("utf-8-sig", "cp1256", "cp1252"):
        try:
            with open(path, encoding=encoding, newline="") as fh:
                return list(csv.DictReader(fh)), []
        except UnicodeDecodeError:
            continue
    raise ValueError(f"cannot decode {path}; save it as 'CSV UTF-8' or as .xlsx")


def import_ground_truth(exp, path: Path, default_verifier: str = "") -> dict:
    rows, natural = read_rows(Path(path))
    report = {"imported": 0, "absent": 0, "empty": 0, "errors": [], "warnings": [],
              "snapshots": set()}
    for i, row in enumerate(rows, start=2):
        sid, field = (row.get("snapshot_id") or "").strip(), (row.get("field") or "").strip()
        if not sid and not field:
            continue
        try:
            snap = Snapshot.objects.get(pk=int(float(sid)), experiment=exp)
            site = site_by_key(exp.config, snap.site_key)
            fdef = next(f for f in site["fields"] if f["name"] == field)
        except (ValueError, Snapshot.DoesNotExist, KeyError, StopIteration):
            report["errors"].append(f"row {i}: unknown snapshot/field ({sid}, {field})")
            continue
        raw = (row.get("value") or "").strip()
        if not raw:
            report["empty"] += 1
            continue
        notes = (row.get("notes") or "").strip()
        if raw.upper() == ABSENT:
            value, notes = "", "absent"
            report["absent"] += 1
        elif fdef["cardinality"] == "many":
            parts = raw.replace("||", "\n").splitlines()
            value = json.dumps([p.strip() for p in parts if p.strip()], ensure_ascii=False)
        elif "||" in raw:
            # Several acceptable forms of one value (e.g. the same instant in two time zones).
            forms = [p.strip() for p in raw.split("||") if p.strip()]
            value = json.dumps({"any": forms}, ensure_ascii=False)
        else:
            value = raw
        GroundTruth.objects.update_or_create(
            snapshot=snap, field_name=field,
            defaults={
                "cardinality": fdef["cardinality"], "value": value, "source": GTSource.HUMAN,
                "verified_by": (row.get("verified_by") or default_verifier or "").strip(),
                "notes": notes,
            },
        )
        report["imported"] += 1
        report["snapshots"].add(snap.pk)

    for row in natural:
        sid = (row.get("snapshot_id") or "").strip()
        answer = (row.get("layout_changed") or "").strip().lower()
        if not sid or answer not in ("yes", "no", "y", "n", "true", "false"):
            continue
        snap = Snapshot.objects.filter(pk=int(float(sid)), experiment=exp).first()
        if snap:
            snap.mutation = {**(snap.mutation or {}), "op": "natural",
                             "layout_changed": answer in ("yes", "y", "true")}
            snap.save(update_fields=["mutation"])

    # Presence checks: typos, and values lost by pruning (fix with prune_keep).
    for sid in sorted(report["snapshots"]):
        snap = Snapshot.objects.get(pk=sid)
        site = site_by_key(exp.config, snap.site_key)
        types = {f["name"]: f["type"] for f in site["fields"]}
        presence = snapstore.compute_field_presence(snap, site["prune_keep"], types)
        for name, p in presence.items():
            if p.get("labelled_absent"):
                continue
            if not p.get("in_raw"):
                report["warnings"].append(
                    f"{snap.site_key}.{name}: value not found in the saved page text — "
                    "check for a typo (it may also live only in an attribute)")
            elif not p.get("in_clean_dom"):
                report["warnings"].append(
                    f"{snap.site_key}.{name}: value is on the page but removed by pruning — "
                    "add its container tag (e.g. aside, nav, header) to the site's prune_keep")
            elif not p.get("in_llm_view"):
                report["warnings"].append(
                    f"{snap.site_key}.{name}: value hidden in the LLM view (long text was "
                    "shortened) — fine for extraction, noted for the paper")
    report["snapshots"] = sorted(report["snapshots"])
    return report
